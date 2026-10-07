"""Export tray-aligned and detailed leaf measurements to an Excel workbook."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from models import GridInfo


def export_excel(output_path, source_image, grid: GridInfo, cell_results):
    output_path = Path(output_path)
    workbook = Workbook()

    # ------------------------------------------------------------------
    # Sheet 1: one Excel cell == one physical tray cell.
    # No header row/column is inserted because that would break the mapping.
    # ------------------------------------------------------------------
    tray_sheet = workbook.active
    tray_sheet.title = "Tray Map"

    for cell in cell_results:
        sheet_cell = tray_sheet.cell(row=cell.row, column=cell.column)

        if cell.status == "NO_SEEDLING":
            value = "NO SEEDLING"
            fill = PatternFill("solid", fgColor="FFE699")
        elif cell.status == "SEEDLING":
            value = (
                f"Leaves: {len(cell.leaves)}\n"
                f"Area: {cell.total_leaf_area_mm2:.2f} mm^2"
            )
            fill = PatternFill("solid", fgColor="C6E0B4")
        else:
            value = cell.status
            fill = PatternFill("solid", fgColor="F4B084")

        if cell.quality_flag != "OK":
            value += "\nREVIEW"

        sheet_cell.value = value
        sheet_cell.fill = fill
        sheet_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet_cell.font = Font(size=9)

    for col in range(1, grid.columns + 1):
        tray_sheet.column_dimensions[get_column_letter(col)].width = 18
    for row in range(1, grid.rows + 1):
        tray_sheet.row_dimensions[row].height = 42

    tray_sheet.freeze_panes = None

    # ------------------------------------------------------------------
    # Sheet 2: detailed leaf-by-leaf measurements.
    # ------------------------------------------------------------------
    detail = workbook.create_sheet("Leaf Measurements")
    headers = [
        "Tray Row",
        "Tray Column",
        "Excel Cell",
        "Cell Status",
        "Cell Quality",
        "Leaf #",
        "Leaf Length (mm)",
        "Leaf Width (mm)",
        "Leaf Area (mm^2)",
        "Leaf Area (px^2)",
        "Mean Leaf Color (HEX)",
        "Leaf Quality",
        "Center X (px)",
        "Center Y (px)",
    ]
    detail.append(headers)

    for header_cell in detail[1]:
        header_cell.font = Font(bold=True)
        header_cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in cell_results:
        if not cell.leaves:
            detail.append([
                cell.row,
                cell.column,
                cell.excel_cell,
                cell.status,
                cell.quality_flag,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
            ])
            continue

        for leaf in cell.leaves:
            detail.append([
                cell.row,
                cell.column,
                cell.excel_cell,
                cell.status,
                cell.quality_flag,
                leaf.leaf_number,
                round(leaf.length_mm, 4),
                round(leaf.width_mm, 4),
                round(leaf.area_mm2, 4),
                round(leaf.area_px, 2),
                leaf.mean_hex,
                leaf.quality_flag,
                round(leaf.center_x_px, 2),
                round(leaf.center_y_px, 2),
            ])

    widths = [11, 13, 12, 14, 14, 8, 18, 17, 18, 18, 22, 16, 14, 14]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    # ------------------------------------------------------------------
    # Sheet 3: calibration and run metadata.
    # ------------------------------------------------------------------
    metadata = workbook.create_sheet("Metadata")
    rows = [
        ("Source image", str(source_image)),
        ("Processed at", datetime.now().isoformat(timespec="seconds")),
        ("Detected rows", grid.rows),
        ("Detected columns", grid.columns),
        ("Detected horizontal pitch (px)", round(grid.pitch_x_px, 4)),
        ("Detected vertical pitch (px)", round(grid.pitch_y_px, 4)),
        ("Assumed physical X pitch (in)", config.CELL_PITCH_X_IN),
        ("Assumed physical Y pitch (in)", config.CELL_PITCH_Y_IN),
        ("X scale (mm/px)", round(grid.mm_per_px_x, 6)),
        ("Y scale (mm/px)", round(grid.mm_per_px_y, 6)),
        ("Green hue range", f"{config.GREEN_H_MIN}..{config.GREEN_H_MAX}"),
        ("Green saturation minimum", config.GREEN_S_MIN),
        ("Green value minimum", config.GREEN_V_MIN),
        ("Excess-green minimum", config.EXCESS_GREEN_MIN),
        ("Cell inset fraction", config.CELL_INSET_FRACTION),
    ]

    for key, value in rows:
        metadata.append([key, value])
    metadata.column_dimensions["A"].width = 34
    metadata.column_dimensions["B"].width = 50
    for cell in metadata["A"]:
        cell.font = Font(bold=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
